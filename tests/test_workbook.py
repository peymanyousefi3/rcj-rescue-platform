import datetime as dt
import io
import json
import zipfile
from pathlib import Path
from zoneinfo import ZoneInfo
import pytest
from django.core.exceptions import ValidationError
from openpyxl import Workbook
from arena.workbook import parse_workbook

def public_workbook(*,mode='Sample'):
    snapshot=json.loads((Path(__file__).resolve().parents[1]/'data/asr26-public-snapshot.json').read_text())
    w=Workbook();s=w.active;s.title='Rescue Simulation 2026'
    for day,col in [(22,2),(23,7),(24,12),(25,17)]:
        s.cell(8,col,f'{day} October');s.cell(10,col,'Time');s.cell(10,col+1,'Programme' if day==22 else 'Rnd')
        for index in range(35):
            total=8*60+30+15*index
            s.cell(index+11,col,dt.time(total//60,total%60))
    first_round=set()
    for item in snapshot['events']:
        when=dt.datetime.fromisoformat(item['starts_at'].replace('Z','+00:00')).astimezone(ZoneInfo('America/Toronto'))
        col={22:2,23:7,24:12,25:17}[when.day]
        row=11+(when.hour*60+when.minute-510)//15
        if item['round_number']:
            n=item['round_number']
            if n not in first_round: s.cell(row,col+1,f'R{n}');first_round.add(n)
            s.cell(row,col+2,item['team_code'])
        else: s.cell(row,col+1,item['title']+(' (TBC)' if not item['confirmed'] else ''))
    s.cell(47,2,'Teams and best-six totals')
    for col,title in [(2,'Team name'),(6,'Code'),(9,'Country'),(11,'School')]: s.cell(48,col,title)
    for i,t in enumerate(snapshot['teams'],49):
        for col,label in [(2,'Team name'),(6,'Code'),(9,'Country'),(11,'School')]:s.cell(i,col,t[label])
    score=w.create_sheet('Score inputs');score.append(['Score inputs · fixed team order']);score.append([])
    score.append(['Mode',mode]);score.append([])
    score.append(['Code','Team']+[f'R{n} raw' for n in range(1,8)])
    scheduled=[t for t in snapshot['teams'] if t['Code']!='105ZZ']
    for t in scheduled: score.append([t['Code'],t['Team name']]+[0]*7)
    score.append([]);score.append(['Code','Team']+[f'R{n} norm.' for n in range(1,8)])
    for t in scheduled: score.append([t['Code'],t['Team name']]+[0]*7)
    stream=io.BytesIO();w.save(stream)
    # Mirror the actual organizer file's missing dimensions / ragged streaming rows.
    result=io.BytesIO()
    import re
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as old,zipfile.ZipFile(result,'w') as new:
        for name in old.namelist():
            data=old.read(name)
            if name.startswith('xl/worksheets/'):
                data=re.sub(rb'<dimension[^>]*/>',b'',data)
            new.writestr(name,data)
    return result.getvalue()

def test_real_layout_contract_includes_five_roster_four_scheduled():
    result=parse_workbook(public_workbook())
    assert len(result['teams'])==5
    assert len(result['rounds'])==7
    assert len(result['events'])==60
    assert {e['team_code'] for e in result['events'] if e['team_code']}=={'101ZZ','102ZZ','103ZZ','104ZZ'}
    assert all(s['raw'] is None and s['normalized'] is None for s in result['scores'])

def test_live_zero_scores_are_retained():
    result=parse_workbook(public_workbook(mode='Live'))
    assert len(result['scores'])==28
    assert all(s['raw']==0 and s['normalized']==0 for s in result['scores'])

def test_unknown_score_mode_fails_closed():
    with pytest.raises(ValidationError): parse_workbook(public_workbook(mode='Unrecognized'))
